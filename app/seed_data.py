"""Validate and import database initialization data from an Excel workbook."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
import re
from typing import TYPE_CHECKING, Any

from openpyxl import load_workbook

if TYPE_CHECKING:
    from sqlalchemy.engine import Connection


class SeedDataError(ValueError):
    """Raised when the seed workbook is missing, malformed, or inconsistent."""


@dataclass(frozen=True)
class SheetSpec:
    required_columns: tuple[str, ...]
    key_column: str


@dataclass(frozen=True)
class SeedWorkbook:
    path: Path
    sheets: dict[str, list[dict[str, Any]]]

    def rows(self, sheet_name: str) -> list[dict[str, Any]]:
        return self.sheets[sheet_name]


@dataclass(frozen=True)
class SeedImportResult:
    row_counts: dict[str, int]

    @property
    def total_rows(self) -> int:
        return sum(self.row_counts.values())


SHEET_SPECS: dict[str, SheetSpec] = {
    "DataSources": SheetSpec(
        ("source_key", "source_name", "source_type", "enabled"), "source_key"
    ),
    "UserTypes": SheetSpec(
        ("user_type_key", "user_type_name"), "user_type_key"
    ),
    "ProductTypes": SheetSpec(
        ("product_type_key", "product_type_name", "max_quantity"),
        "product_type_key",
    ),
    "Assets": SheetSpec(
        ("asset_key", "asset_name", "cmms_asset_id", "dc_asset_id"), "asset_key"
    ),
    "MeasurementTypes": SheetSpec(
        (
            "measurement_type_key",
            "measurement_type_name",
            "unit",
            "unit_name",
            "unit_symbol",
        ),
        "measurement_type_key",
    ),
    "SensorTypes": SheetSpec(
        (
            "sensor_type_key",
            "type_name",
            "max_value",
            "min_value",
            "measurement_type_key",
            "accuracy",
        ),
        "sensor_type_key",
    ),
    "Sensors": SheetSpec(
        (
            "sensor_key",
            "sensor_name",
            "sensor_type_key",
            "asset_key",
            "measurement_frequency",
            "ranges_key",
            "metric_function_id",
            "chart_aggregation_method",
        ),
        "sensor_key",
    ),
    "ProcessConfigurations": SheetSpec(
        (
            "configuration_key",
            "configuration_name",
            "valid_from",
            "valid_to",
            "is_active",
        ),
        "configuration_key",
    ),
    "ProcessSteps": SheetSpec(
        ("process_step_key", "process_step_name", "processing_time"),
        "process_step_key",
    ),
    "AssetProcessSteps": SheetSpec(
        ("configuration_key", "asset_key", "process_step_key"),
        "configuration_key",
    ),
    "Routing": SheetSpec(
        ("configuration_key", "process_step_key", "next_process_step_key"),
        "configuration_key",
    ),
}

_KEY_PATTERN = re.compile(r"^[a-z0-9][a-z0-9_-]{0,63}$")
_AGGREGATION_METHODS = {"average", "latest", "minimum", "maximum"}
_SOURCE_TYPES = {"physical", "simulation"}


def _empty_to_none(value: Any) -> Any:
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def _location(sheet: str, row: dict[str, Any], column: str) -> str:
    return f"{sheet}!{column} (row {row['_excel_row']})"


def _required_text(sheet: str, row: dict[str, Any], column: str) -> str:
    value = _empty_to_none(row.get(column))
    if value is None:
        raise SeedDataError(f"Missing required value at {_location(sheet, row, column)}")
    return str(value).strip()


def _optional_text(value: Any) -> str | None:
    value = _empty_to_none(value)
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _number(
    sheet: str,
    row: dict[str, Any],
    column: str,
    *,
    required: bool = False,
) -> float | None:
    value = _empty_to_none(row.get(column))
    if value is None:
        if required:
            raise SeedDataError(
                f"Missing required value at {_location(sheet, row, column)}"
            )
        return None
    if isinstance(value, bool):
        raise SeedDataError(f"Expected a number at {_location(sheet, row, column)}")
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise SeedDataError(
            f"Expected a number at {_location(sheet, row, column)}: {value!r}"
        ) from exc


def _integer(
    sheet: str,
    row: dict[str, Any],
    column: str,
    *,
    required: bool = False,
) -> int | None:
    value = _number(sheet, row, column, required=required)
    if value is None:
        return None
    if not value.is_integer():
        raise SeedDataError(f"Expected an integer at {_location(sheet, row, column)}")
    return int(value)


def _boolean(sheet: str, row: dict[str, Any], column: str) -> bool:
    value = _empty_to_none(row.get(column))
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value)
    if isinstance(value, str):
        normalized = value.casefold()
        if normalized in {"true", "yes", "1"}:
            return True
        if normalized in {"false", "no", "0"}:
            return False
    raise SeedDataError(f"Expected TRUE/FALSE at {_location(sheet, row, column)}")


def _optional_datetime(
    sheet: str, row: dict[str, Any], column: str
) -> datetime | None:
    value = _empty_to_none(row.get(column))
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value).replace(tzinfo=None)
        except ValueError as exc:
            raise SeedDataError(
                f"Expected an ISO date/time at {_location(sheet, row, column)}"
            ) from exc
    raise SeedDataError(f"Expected a date/time at {_location(sheet, row, column)}")


def _read_sheet(worksheet: Any, spec: SheetSpec) -> list[dict[str, Any]]:
    header_row: int | None = None
    headers: list[str | None] = []
    required = set(spec.required_columns)

    for row_number in range(1, min(worksheet.max_row, 25) + 1):
        candidate = [
            _optional_text(worksheet.cell(row=row_number, column=column).value)
            for column in range(1, worksheet.max_column + 1)
        ]
        if required.issubset({value for value in candidate if value is not None}):
            header_row = row_number
            headers = candidate
            break

    if header_row is None:
        raise SeedDataError(
            f"Sheet {worksheet.title!r} does not contain the required headers: "
            + ", ".join(spec.required_columns)
        )

    nonblank_headers = [header for header in headers if header is not None]
    duplicates = sorted(
        {header for header in nonblank_headers if nonblank_headers.count(header) > 1}
    )
    if duplicates:
        raise SeedDataError(
            f"Duplicate headers on sheet {worksheet.title!r}: {', '.join(duplicates)}"
        )

    rows: list[dict[str, Any]] = []
    for row_number in range(header_row + 1, worksheet.max_row + 1):
        values = [
            _empty_to_none(worksheet.cell(row=row_number, column=column).value)
            for column in range(1, worksheet.max_column + 1)
        ]
        if all(value is None for value in values):
            continue
        row = {
            header: values[index]
            for index, header in enumerate(headers)
            if header is not None
        }
        row["_excel_row"] = row_number
        rows.append(row)

    return rows


def _keys(seed: SeedWorkbook, sheet: str, column: str) -> set[str]:
    return {_required_text(sheet, row, column) for row in seed.rows(sheet)}


def _validate_unique_keys(seed: SeedWorkbook) -> None:
    for sheet, spec in SHEET_SPECS.items():
        if sheet in {"AssetProcessSteps", "Routing"}:
            continue
        seen: dict[str, int] = {}
        for row in seed.rows(sheet):
            key = _required_text(sheet, row, spec.key_column)
            if not _KEY_PATTERN.fullmatch(key):
                raise SeedDataError(
                    f"Invalid stable key at {_location(sheet, row, spec.key_column)}: "
                    f"{key!r}; use lowercase letters, numbers, '_' or '-'"
                )
            if key in seen:
                raise SeedDataError(
                    f"Duplicate key {key!r} in {sheet}, rows {seen[key]} and "
                    f"{row['_excel_row']}"
                )
            seen[key] = row["_excel_row"]


def _validate_optional_unique(
    sheet: str, rows: list[dict[str, Any]], column: str
) -> None:
    seen: dict[Any, int] = {}
    for row in rows:
        value = _empty_to_none(row.get(column))
        if value is None:
            continue
        if value in seen:
            raise SeedDataError(
                f"Duplicate {column} {value!r} in {sheet}, rows {seen[value]} and "
                f"{row['_excel_row']}"
            )
        seen[value] = row["_excel_row"]


def _validate_reference(
    sheet: str,
    rows: list[dict[str, Any]],
    column: str,
    allowed: set[str],
) -> None:
    for row in rows:
        value = _required_text(sheet, row, column)
        if value not in allowed:
            raise SeedDataError(
                f"Unknown reference {value!r} at {_location(sheet, row, column)}"
            )


def _validate_seed(seed: SeedWorkbook) -> None:
    _validate_unique_keys(seed)

    for row in seed.rows("Assets"):
        _required_text("Assets", row, "asset_name")
        _integer("Assets", row, "cmms_asset_id")
        _optional_text(row.get("dc_asset_id"))
    _validate_optional_unique("Assets", seed.rows("Assets"), "cmms_asset_id")
    _validate_optional_unique("Assets", seed.rows("Assets"), "dc_asset_id")

    for row in seed.rows("MeasurementTypes"):
        _required_text("MeasurementTypes", row, "measurement_type_name")
        _required_text("MeasurementTypes", row, "unit")

    for row in seed.rows("UserTypes"):
        _required_text("UserTypes", row, "user_type_name")

    for row in seed.rows("ProductTypes"):
        _required_text("ProductTypes", row, "product_type_name")

    asset_keys = _keys(seed, "Assets", "asset_key")
    measurement_type_keys = _keys(
        seed, "MeasurementTypes", "measurement_type_key"
    )
    sensor_type_keys = _keys(seed, "SensorTypes", "sensor_type_key")
    configuration_keys = _keys(
        seed, "ProcessConfigurations", "configuration_key"
    )
    process_step_keys = _keys(seed, "ProcessSteps", "process_step_key")

    _validate_reference(
        "SensorTypes",
        seed.rows("SensorTypes"),
        "measurement_type_key",
        measurement_type_keys,
    )
    _validate_reference(
        "Sensors", seed.rows("Sensors"), "sensor_type_key", sensor_type_keys
    )
    _validate_reference("Sensors", seed.rows("Sensors"), "asset_key", asset_keys)
    for row in seed.rows("Sensors"):
        _required_text("Sensors", row, "sensor_name")
        frequency = _number(
            "Sensors", row, "measurement_frequency", required=True
        )
        if frequency is not None and frequency <= 0:
            raise SeedDataError(
                f"measurement_frequency must be positive at "
                f"{_location('Sensors', row, 'measurement_frequency')}"
            )
        aggregation = _required_text(
            "Sensors", row, "chart_aggregation_method"
        )
        if aggregation not in _AGGREGATION_METHODS:
            raise SeedDataError(
                f"Invalid chart aggregation method {aggregation!r} at "
                f"{_location('Sensors', row, 'chart_aggregation_method')}"
            )
        if _optional_text(row.get("ranges_key")) is not None:
            raise SeedDataError(
                f"ranges_key is reserved but not importable until a Ranges sheet is "
                f"defined; clear {_location('Sensors', row, 'ranges_key')}"
            )

    for row in seed.rows("SensorTypes"):
        _required_text("SensorTypes", row, "type_name")
        minimum = _number("SensorTypes", row, "min_value")
        maximum = _number("SensorTypes", row, "max_value")
        _number("SensorTypes", row, "accuracy")
        if minimum is not None and maximum is not None and minimum > maximum:
            raise SeedDataError(
                f"min_value exceeds max_value in SensorTypes row {row['_excel_row']}"
            )

    active_count = 0
    for row in seed.rows("ProcessConfigurations"):
        _required_text("ProcessConfigurations", row, "configuration_name")
        active_count += int(_boolean("ProcessConfigurations", row, "is_active"))
        valid_from = _optional_datetime(
            "ProcessConfigurations", row, "valid_from"
        )
        valid_to = _optional_datetime("ProcessConfigurations", row, "valid_to")
        if valid_from is not None and valid_to is not None and valid_to <= valid_from:
            raise SeedDataError(
                f"valid_to must be later than valid_from in ProcessConfigurations "
                f"row {row['_excel_row']}"
            )
    if active_count != 1:
        raise SeedDataError(
            "ProcessConfigurations must contain exactly one active configuration"
        )

    for row in seed.rows("ProcessSteps"):
        _required_text("ProcessSteps", row, "process_step_name")
        processing_time = _number(
            "ProcessSteps", row, "processing_time", required=True
        )
        if processing_time is not None and processing_time < 0:
            raise SeedDataError(
                f"processing_time must not be negative in ProcessSteps row "
                f"{row['_excel_row']}"
            )

    for row in seed.rows("ProductTypes"):
        maximum = _integer("ProductTypes", row, "max_quantity")
        if maximum is not None and maximum <= 0:
            raise SeedDataError(
                f"max_quantity must be positive in ProductTypes row "
                f"{row['_excel_row']}"
            )

    for row in seed.rows("DataSources"):
        _required_text("DataSources", row, "source_name")
        source_type = _required_text("DataSources", row, "source_type")
        if source_type not in _SOURCE_TYPES:
            raise SeedDataError(
                f"Invalid source_type {source_type!r} at "
                f"{_location('DataSources', row, 'source_type')}"
            )
        _boolean("DataSources", row, "enabled")
    if "real" not in _keys(seed, "DataSources", "source_key"):
        raise SeedDataError("DataSources must contain the 'real' source_key")

    for relation_sheet in ("AssetProcessSteps", "Routing"):
        _validate_reference(
            relation_sheet,
            seed.rows(relation_sheet),
            "configuration_key",
            configuration_keys,
        )
        _validate_reference(
            relation_sheet,
            seed.rows(relation_sheet),
            "process_step_key",
            process_step_keys,
        )
    _validate_reference(
        "AssetProcessSteps",
        seed.rows("AssetProcessSteps"),
        "asset_key",
        asset_keys,
    )
    _validate_reference(
        "Routing",
        seed.rows("Routing"),
        "next_process_step_key",
        process_step_keys,
    )

    for sheet, columns in (
        ("AssetProcessSteps", ("configuration_key", "asset_key", "process_step_key")),
        ("Routing", ("configuration_key", "process_step_key", "next_process_step_key")),
    ):
        seen: dict[tuple[str, ...], int] = {}
        for row in seed.rows(sheet):
            key = tuple(_required_text(sheet, row, column) for column in columns)
            if sheet == "Routing" and key[1] == key[2]:
                raise SeedDataError(f"Routing self-reference in row {row['_excel_row']}")
            if key in seen:
                raise SeedDataError(
                    f"Duplicate relation in {sheet}, rows {seen[key]} and "
                    f"{row['_excel_row']}"
                )
            seen[key] = row["_excel_row"]


def load_seed_workbook(path: str | Path) -> SeedWorkbook:
    """Read and fully validate the workbook without touching the database."""

    workbook_path = Path(path).expanduser().resolve()
    if not workbook_path.is_file():
        raise SeedDataError(f"Seed workbook does not exist: {workbook_path}")

    try:
        workbook = load_workbook(workbook_path, read_only=False, data_only=True)
    except Exception as exc:
        raise SeedDataError(f"Cannot read seed workbook {workbook_path}: {exc}") from exc

    try:
        missing_sheets = set(SHEET_SPECS) - set(workbook.sheetnames)
        if missing_sheets:
            raise SeedDataError(
                "Seed workbook is missing sheets: " + ", ".join(sorted(missing_sheets))
            )
        sheets = {
            sheet_name: _read_sheet(workbook[sheet_name], spec)
            for sheet_name, spec in SHEET_SPECS.items()
        }
    finally:
        workbook.close()

    seed = SeedWorkbook(path=workbook_path, sheets=sheets)
    _validate_seed(seed)
    return seed


def _claim_legacy_row(
    connection: Connection,
    *,
    table: str,
    id_column: str,
    key_column: str,
    key: str,
    natural_predicate: str,
    parameters: dict[str, Any],
) -> None:
    """Attach a stable key to one pre-Excel row created by the former SQL seed."""

    from sqlalchemy import text

    row_id = connection.execute(
        text(
            f"SELECT {id_column} FROM public.{table} "
            f"WHERE {key_column} IS NULL AND {natural_predicate} "
            f"ORDER BY {id_column} LIMIT 1"
        ),
        parameters,
    ).scalar_one_or_none()
    if row_id is not None:
        connection.execute(
            text(
                f"UPDATE public.{table} SET {key_column} = :stable_key "
                f"WHERE {id_column} = :row_id"
            ),
            {"stable_key": key, "row_id": row_id},
        )


def _upsert_returning_id(
    connection: Connection,
    statement: str,
    parameters: dict[str, Any],
) -> int:
    from sqlalchemy import text

    return int(connection.execute(text(statement), parameters).scalar_one())


def import_seed_workbook(connection: Connection, seed: SeedWorkbook) -> SeedImportResult:
    """Upsert a validated workbook using its stable keys in one caller transaction."""

    from sqlalchemy import text

    row_counts = {sheet: len(rows) for sheet, rows in seed.sheets.items()}

    data_source_ids: dict[str, int] = {}
    for row in seed.rows("DataSources"):
        key = _required_text("DataSources", row, "source_key")
        data_source_ids[key] = _upsert_returning_id(
            connection,
            """
            INSERT INTO public.data_sources
                (source_key, source_name, source_type, enabled)
            VALUES (:key, :name, :source_type, :enabled)
            ON CONFLICT (source_key) DO UPDATE SET
                source_name = EXCLUDED.source_name,
                source_type = EXCLUDED.source_type,
                enabled = EXCLUDED.enabled
            RETURNING data_source_id
            """,
            {
                "key": key,
                "name": _required_text("DataSources", row, "source_name"),
                "source_type": _required_text("DataSources", row, "source_type"),
                "enabled": _boolean("DataSources", row, "enabled"),
            },
        )

    asset_ids: dict[str, int] = {}
    for row in seed.rows("Assets"):
        key = _required_text("Assets", row, "asset_key")
        name = _required_text("Assets", row, "asset_name")
        _claim_legacy_row(
            connection,
            table="assets",
            id_column="asset_id",
            key_column="asset_key",
            key=key,
            natural_predicate="asset_name = :name",
            parameters={"name": name},
        )
        asset_ids[key] = _upsert_returning_id(
            connection,
            """
            INSERT INTO public.assets
                (asset_key, asset_name, cmms_asset_id, dc_asset_id)
            VALUES (:key, :name, :cmms_asset_id, :dc_asset_id)
            ON CONFLICT (asset_key) DO UPDATE SET
                asset_name = EXCLUDED.asset_name,
                cmms_asset_id = COALESCE(EXCLUDED.cmms_asset_id, assets.cmms_asset_id),
                dc_asset_id = COALESCE(EXCLUDED.dc_asset_id, assets.dc_asset_id)
            RETURNING asset_id
            """,
            {
                "key": key,
                "name": name,
                "cmms_asset_id": _integer("Assets", row, "cmms_asset_id"),
                "dc_asset_id": _optional_text(row.get("dc_asset_id")),
            },
        )

    measurement_type_ids: dict[str, int] = {}
    for row in seed.rows("MeasurementTypes"):
        key = _required_text("MeasurementTypes", row, "measurement_type_key")
        name = _required_text("MeasurementTypes", row, "measurement_type_name")
        unit = _required_text("MeasurementTypes", row, "unit")
        _claim_legacy_row(
            connection,
            table="measurement_types",
            id_column="measurement_type_id",
            key_column="measurement_type_key",
            key=key,
            natural_predicate="measurement_type_name = :name AND unit = :unit",
            parameters={"name": name, "unit": unit},
        )
        measurement_type_ids[key] = _upsert_returning_id(
            connection,
            """
            INSERT INTO public.measurement_types
                (measurement_type_key, measurement_type_name, unit, unit_name, unit_symbol)
            VALUES (:key, :name, :unit, :unit_name, :unit_symbol)
            ON CONFLICT (measurement_type_key) DO UPDATE SET
                measurement_type_name = EXCLUDED.measurement_type_name,
                unit = EXCLUDED.unit,
                unit_name = COALESCE(EXCLUDED.unit_name, measurement_types.unit_name),
                unit_symbol = COALESCE(EXCLUDED.unit_symbol, measurement_types.unit_symbol)
            RETURNING measurement_type_id
            """,
            {
                "key": key,
                "name": name,
                "unit": unit,
                "unit_name": _optional_text(row.get("unit_name")),
                "unit_symbol": _optional_text(row.get("unit_symbol")),
            },
        )

    sensor_type_ids: dict[str, int] = {}
    for row in seed.rows("SensorTypes"):
        key = _required_text("SensorTypes", row, "sensor_type_key")
        name = _required_text("SensorTypes", row, "type_name")
        measurement_type_id = measurement_type_ids[
            _required_text("SensorTypes", row, "measurement_type_key")
        ]
        _claim_legacy_row(
            connection,
            table="sensor_types",
            id_column="type_id",
            key_column="sensor_type_key",
            key=key,
            natural_predicate="type_name = :name AND measurement_type_id = :measurement_type_id",
            parameters={"name": name, "measurement_type_id": measurement_type_id},
        )
        sensor_type_ids[key] = _upsert_returning_id(
            connection,
            """
            INSERT INTO public.sensor_types
                (sensor_type_key, type_name, max_value, min_value,
                 measurement_type_id, accuracy)
            VALUES (:key, :name, :max_value, :min_value,
                    :measurement_type_id, :accuracy)
            ON CONFLICT (sensor_type_key) DO UPDATE SET
                type_name = EXCLUDED.type_name,
                max_value = EXCLUDED.max_value,
                min_value = EXCLUDED.min_value,
                measurement_type_id = EXCLUDED.measurement_type_id,
                accuracy = COALESCE(EXCLUDED.accuracy, sensor_types.accuracy)
            RETURNING type_id
            """,
            {
                "key": key,
                "name": name,
                "max_value": _number("SensorTypes", row, "max_value"),
                "min_value": _number("SensorTypes", row, "min_value"),
                "measurement_type_id": measurement_type_id,
                "accuracy": _number("SensorTypes", row, "accuracy"),
            },
        )

    sensor_ids: dict[str, int] = {}
    for row in seed.rows("Sensors"):
        key = _required_text("Sensors", row, "sensor_key")
        name = _required_text("Sensors", row, "sensor_name")
        type_id = sensor_type_ids[
            _required_text("Sensors", row, "sensor_type_key")
        ]
        asset_id = asset_ids[_required_text("Sensors", row, "asset_key")]
        _claim_legacy_row(
            connection,
            table="sensors",
            id_column="sensor_id",
            key_column="sensor_key",
            key=key,
            natural_predicate=(
                "sensor_name = :name AND type_id = :type_id AND asset_id = :asset_id"
            ),
            parameters={"name": name, "type_id": type_id, "asset_id": asset_id},
        )
        sensor_ids[key] = _upsert_returning_id(
            connection,
            """
            INSERT INTO public.sensors
                (sensor_key, sensor_name, measurement_frequency, ranges_id,
                 type_id, asset_id, metric_function_id, chart_aggregation_method)
            VALUES (:key, :name, :measurement_frequency, NULL,
                    :type_id, :asset_id, :metric_function_id, :aggregation)
            ON CONFLICT (sensor_key) DO UPDATE SET
                sensor_name = EXCLUDED.sensor_name,
                measurement_frequency = EXCLUDED.measurement_frequency,
                type_id = EXCLUDED.type_id,
                asset_id = EXCLUDED.asset_id,
                metric_function_id = COALESCE(
                    EXCLUDED.metric_function_id,
                    sensors.metric_function_id
                ),
                chart_aggregation_method = EXCLUDED.chart_aggregation_method
            RETURNING sensor_id
            """,
            {
                "key": key,
                "name": name,
                "measurement_frequency": _number(
                    "Sensors", row, "measurement_frequency", required=True
                ),
                "type_id": type_id,
                "asset_id": asset_id,
                "metric_function_id": _optional_text(row.get("metric_function_id")),
                "aggregation": _required_text(
                    "Sensors", row, "chart_aggregation_method"
                ),
            },
        )

    product_type_ids: dict[str, int] = {}
    for row in seed.rows("ProductTypes"):
        key = _required_text("ProductTypes", row, "product_type_key")
        name = _required_text("ProductTypes", row, "product_type_name")
        _claim_legacy_row(
            connection,
            table="product_types",
            id_column="product_type_id",
            key_column="product_type_key",
            key=key,
            natural_predicate="product_type_name = :name",
            parameters={"name": name},
        )
        product_type_ids[key] = _upsert_returning_id(
            connection,
            """
            INSERT INTO public.product_types
                (product_type_key, product_type_name, max_quantity)
            VALUES (:key, :name, :max_quantity)
            ON CONFLICT (product_type_key) DO UPDATE SET
                product_type_name = EXCLUDED.product_type_name,
                max_quantity = EXCLUDED.max_quantity
            RETURNING product_type_id
            """,
            {
                "key": key,
                "name": name,
                "max_quantity": _integer("ProductTypes", row, "max_quantity"),
            },
        )

    user_type_ids: dict[str, int] = {}
    for row in seed.rows("UserTypes"):
        key = _required_text("UserTypes", row, "user_type_key")
        name = _required_text("UserTypes", row, "user_type_name")
        _claim_legacy_row(
            connection,
            table="user_types",
            id_column="user_type_id",
            key_column="user_type_key",
            key=key,
            natural_predicate="user_type_name = :name",
            parameters={"name": name},
        )
        user_type_ids[key] = _upsert_returning_id(
            connection,
            """
            INSERT INTO public.user_types
                (user_type_key, user_type_name)
            VALUES (:key, :name)
            ON CONFLICT (user_type_key) DO UPDATE SET
                user_type_name = EXCLUDED.user_type_name
            RETURNING user_type_id
            """,
            {"key": key, "name": name},
        )

    configuration_ids: dict[str, int] = {}
    connection.execute(text("UPDATE public.process_configurations SET is_active = FALSE"))
    for row in seed.rows("ProcessConfigurations"):
        key = _required_text(
            "ProcessConfigurations", row, "configuration_key"
        )
        name = _required_text(
            "ProcessConfigurations", row, "configuration_name"
        )
        _claim_legacy_row(
            connection,
            table="process_configurations",
            id_column="process_configuration_id",
            key_column="configuration_key",
            key=key,
            natural_predicate="configuration_name = :name",
            parameters={"name": name},
        )
        configuration_ids[key] = _upsert_returning_id(
            connection,
            """
            INSERT INTO public.process_configurations
                (configuration_key, configuration_name, valid_from, valid_to, is_active)
            VALUES (:key, :name, COALESCE(:valid_from, CURRENT_TIMESTAMP),
                    :valid_to, :is_active)
            ON CONFLICT (configuration_key) DO UPDATE SET
                configuration_name = EXCLUDED.configuration_name,
                valid_from = COALESCE(:valid_from, process_configurations.valid_from),
                valid_to = EXCLUDED.valid_to,
                is_active = EXCLUDED.is_active
            RETURNING process_configuration_id
            """,
            {
                "key": key,
                "name": name,
                "valid_from": _optional_datetime(
                    "ProcessConfigurations", row, "valid_from"
                ),
                "valid_to": _optional_datetime(
                    "ProcessConfigurations", row, "valid_to"
                ),
                "is_active": _boolean(
                    "ProcessConfigurations", row, "is_active"
                ),
            },
        )

    process_step_ids: dict[str, int] = {}
    for row in seed.rows("ProcessSteps"):
        key = _required_text("ProcessSteps", row, "process_step_key")
        name = _required_text("ProcessSteps", row, "process_step_name")
        processing_time = _number(
            "ProcessSteps", row, "processing_time", required=True
        )
        _claim_legacy_row(
            connection,
            table="process_steps",
            id_column="process_step_id",
            key_column="process_step_key",
            key=key,
            natural_predicate=(
                "process_step_name = :name AND processing_time = :processing_time"
            ),
            parameters={"name": name, "processing_time": processing_time},
        )
        process_step_ids[key] = _upsert_returning_id(
            connection,
            """
            INSERT INTO public.process_steps
                (process_step_key, process_step_name, processing_time)
            VALUES (:key, :name, :processing_time)
            ON CONFLICT (process_step_key) DO UPDATE SET
                process_step_name = EXCLUDED.process_step_name,
                processing_time = EXCLUDED.processing_time
            RETURNING process_step_id
            """,
            {"key": key, "name": name, "processing_time": processing_time},
        )

    for configuration_id in configuration_ids.values():
        connection.execute(
            text(
                "DELETE FROM public.routing "
                "WHERE process_configuration_id = :configuration_id"
            ),
            {"configuration_id": configuration_id},
        )
        connection.execute(
            text(
                "DELETE FROM public.asset_process_steps "
                "WHERE process_configuration_id = :configuration_id"
            ),
            {"configuration_id": configuration_id},
        )

    relation_insert = text(
        """
        INSERT INTO public.asset_process_steps
            (process_configuration_id, asset_id, process_step_id)
        VALUES (:configuration_id, :asset_id, :process_step_id)
        ON CONFLICT (process_configuration_id, asset_id, process_step_id) DO NOTHING
        """
    )
    for row in seed.rows("AssetProcessSteps"):
        connection.execute(
            relation_insert,
            {
                "configuration_id": configuration_ids[
                    _required_text("AssetProcessSteps", row, "configuration_key")
                ],
                "asset_id": asset_ids[
                    _required_text("AssetProcessSteps", row, "asset_key")
                ],
                "process_step_id": process_step_ids[
                    _required_text("AssetProcessSteps", row, "process_step_key")
                ],
            },
        )

    routing_insert = text(
        """
        INSERT INTO public.routing
            (process_configuration_id, process_step_id, next_process_step_id)
        VALUES (:configuration_id, :process_step_id, :next_process_step_id)
        ON CONFLICT
            (process_configuration_id, process_step_id, next_process_step_id)
        DO NOTHING
        """
    )
    for row in seed.rows("Routing"):
        connection.execute(
            routing_insert,
            {
                "configuration_id": configuration_ids[
                    _required_text("Routing", row, "configuration_key")
                ],
                "process_step_id": process_step_ids[
                    _required_text("Routing", row, "process_step_key")
                ],
                "next_process_step_id": process_step_ids[
                    _required_text("Routing", row, "next_process_step_key")
                ],
            },
        )

    # Keep the existing default used by Node-RED and the DataCollector aligned
    # with the workbook's physical source.
    real_source_id = data_source_ids["real"]
    connection.execute(
        text(
            "ALTER TABLE public.measurements ALTER COLUMN data_source_id "
            f"SET DEFAULT {real_source_id}"
        )
    )

    return SeedImportResult(row_counts=row_counts)
